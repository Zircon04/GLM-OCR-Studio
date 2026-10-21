#!/usr/bin/env python3
"""Create local HTML, DOCX, searchable PDF and ZIP exports for OCR output."""

from __future__ import annotations

import argparse
import html
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path
from urllib.parse import quote, unquote, urlsplit

from searchable_pdf import make_searchable_pdf

EVENT_PREFIX = "GLMOCR_EVENT "
RECONSTRUCT_SCRIPT = Path(__file__).with_name("reconstruct_pdf.py")


def emit(kind: str, **payload) -> None:
    print(EVENT_PREFIX + json.dumps({"type": kind, **payload}, ensure_ascii=False), flush=True)


def find_file(root: Path, suffix: str, exclude: str = "") -> Path | None:
    files = [p for p in root.rglob(f"*{suffix}") if not exclude or exclude not in p.name]
    return sorted(files, key=lambda p: len(p.parts))[0] if files else None


_LITERAL = re.compile(r"(```.*?```|\$\$.*?\$\$|`[^`\n]*`|\$[^$\n]+\$)", re.DOTALL)
_MD_IMG = re.compile(r"(!?\[[^\]]*\]\()([^)\s]+)(\))")
_HTML_IMG = re.compile(r"(<img\b[^>]*\bsrc=[\"'])([^\"']+)([\"'])", re.IGNORECASE)
CLI_EXPORTS = {"md", "split_md", "searchable_pdf"}
INTERNAL_EXPORTS = CLI_EXPORTS | {"html", "docx", "clean_pdf", "zip"}


def _rewrite_piece(piece: str, source_dir: Path, dest_dir: Path) -> str:
    def retarget(raw: str) -> str:
        parts = urlsplit(raw)
        if parts.scheme in {"http", "https", "data", "mailto"} or raw.startswith("#"):
            return raw
        decoded = unquote(parts.path)
        if not decoded or decoded.startswith("/"):
            return raw
        target = (source_dir / decoded).resolve()
        relative = Path(os_rel(target, dest_dir)).as_posix()
        return quote(relative, safe="/-._~") + (f"?{parts.query}" if parts.query else "")

    def md(match: re.Match[str]) -> str:
        return match.group(1) + retarget(match.group(2)) + match.group(3)

    def html_img(match: re.Match[str]) -> str:
        return match.group(1) + retarget(match.group(2)) + match.group(3)

    return _HTML_IMG.sub(html_img, _MD_IMG.sub(md, piece))


def os_rel(target: Path, start: Path) -> str:
    import os
    return os.path.relpath(target, start)


def rewrite_image_links(markdown: str, source_dir: Path, dest_file: Path) -> str:
    """Пути картинок считаются от файла, куда пишется Markdown.

    Формулы и код остаются как есть. Исходные ссылки ищутся от папки книги.
    """
    parts = _LITERAL.split(markdown)
    dest_dir = dest_file.parent
    return "".join(
        part if index % 2 else _rewrite_piece(part, source_dir, dest_dir)
        for index, part in enumerate(parts)
    )


def markdown_to_html(markdown: str, base: str) -> str:
    try:
        import markdown as md
        body = md.markdown(markdown, extensions=["tables", "fenced_code", "sane_lists"])
        prefix = "../" + quote(base) + "/"
        body = body.replace('src="imgs/', f'src="{prefix}imgs/')
        body = body.replace('src="imgs_enhanced/', f'src="{prefix}imgs_enhanced/')
    except Exception:
        body = "<pre>" + html.escape(markdown) + "</pre>"
    return f"""<!doctype html><html lang="ru"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(base)}</title><style>
:root{{color-scheme:light dark}}body{{font:16px -apple-system,BlinkMacSystemFont,sans-serif;line-height:1.55;max-width:920px;margin:48px auto;padding:0 28px;color:CanvasText;background:Canvas}}
h1,h2,h3{{line-height:1.2}}img{{max-width:100%;height:auto;border-radius:10px}}table{{border-collapse:collapse;width:100%;overflow:auto}}td,th{{border:1px solid color-mix(in srgb,CanvasText 20%,transparent);padding:8px}}pre,code{{font-family:ui-monospace,SFMono-Regular,monospace}}pre{{overflow:auto;padding:14px;border-radius:10px;background:color-mix(in srgb,CanvasText 7%,Canvas)}}
</style></head><body>{body}</body></html>"""


def add_markdown_to_docx(markdown: str, output: Path, title: str, asset_root: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_SECTION
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    section.header_distance = section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8), ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = styles[name]
        style.font.name = "Arial"; style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
    header = section.header.paragraphs[0]
    header.text = "GLM OCR Local · распознанный документ"
    header.style = styles["Header"]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Локальная обработка MLX/Metal")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run(title)
    run.bold = True; run.font.name = "Arial"; run.font.size = Pt(22); run.font.color.rgb = RGBColor.from_string("0B2545")

    in_code = False
    code_buffer: list[str] = []
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph("\n".join(code_buffer))
                p.style = styles["No Spacing"]
                for r in p.runs: r.font.name = "Menlo"; r.font.size = Pt(8.5)
                code_buffer = []
            in_code = not in_code
            continue
        if in_code:
            code_buffer.append(line); continue
        if not line:
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 3")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 2")
        elif line.startswith("# "):
            doc.add_paragraph(line[2:], style="Heading 1")
        elif re.match(r"^[-*+]\s+", line):
            doc.add_paragraph(re.sub(r"^[-*+]\s+", "", line), style="List Bullet")
        elif re.match(r"^\d+[.)]\s+", line):
            doc.add_paragraph(re.sub(r"^\d+[.)]\s+", "", line), style="List Number")
        elif line.startswith("!["):
            match = re.search(r"\]\(([^)]+)\)", line)
            if match:
                image_path = asset_root / match.group(1)
                if image_path.exists():
                    try: doc.add_picture(str(image_path), width=Inches(5.8))
                    except Exception: pass
        elif line.startswith("<table"):
            p = doc.add_paragraph()
            p.add_run(re.sub(r"<[^>]+>", " ", line)).font.size = Pt(9)
        else:
            p = doc.add_paragraph()
            fragments = re.split(r"(`[^`]+`|\$\$.*?\$\$|\$[^$]+\$)", line)
            for fragment in fragments:
                r = p.add_run(fragment)
                if fragment.startswith("`") or fragment.startswith("$"):
                    r.font.name = "Cambria Math" if fragment.startswith("$") else "Menlo"
                    r.font.size = Pt(9.5)
    doc.core_properties.title = title
    doc.core_properties.subject = "Результат локального OCR"
    doc.core_properties.author = "GLM OCR Local"
    doc.save(output)


def make_split_md(markdown: str, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    lines = markdown.splitlines(keepends=True)
    heading_prefix = '# '
    current_section = []
    section_title = "00_Введение"
    part_counter = 1

    def save_section():
        nonlocal part_counter
        if not current_section:
            return
        
        clean_title = re.sub(r'[\\/*?:"<>|]', "", section_title).strip()
        clean_title = clean_title.replace(" ", "_")
        if not clean_title:
            clean_title = f"part_{part_counter}"
            
        filename = f"{part_counter:02d}_{clean_title}.md"
        filepath = output_dir / filename
        
        filepath.write_text("".join(current_section), encoding='utf-8')
        part_counter += 1

    for line in lines:
        if line.startswith(heading_prefix):
            if current_section:
                save_section()
            section_title = line[len(heading_prefix):].strip()
            current_section = [line]
        else:
            current_section.append(line)

    if current_section:
        save_section()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--result", required=True)
    parser.add_argument("--source", required=True)
    parser.add_argument("--output-root", required=True)
    parser.add_argument("--formats", default="md,split_md,searchable_pdf")
    parser.add_argument("--downsample-dpi", type=int, default=None, help="DPI картинок страницы. Пусто — не менять размер.")
    parser.add_argument("--jpeg-quality", type=int, default=0, help="JPEG 1–95. 0 — не пережимать. По умолчанию 0.")
    parser.add_argument("--force-grayscale", action="store_true", help="Перевести картинки страниц в оттенки серого.")
    parser.add_argument("--overlay-crops", action="store_true", help="Наложить вырезанные моделью фигуры поверх скана.")
    args = parser.parse_args()
    result = Path(args.result)
    source = Path(args.source)
    output_root = Path(args.output_root)
    requested = [item.strip() for item in args.formats.split(",") if item.strip()]
    if not requested:
        emit("export_error", message="Список форматов пуст")
        return 2
    unknown = [item for item in requested if item not in INTERNAL_EXPORTS]
    if unknown:
        emit("export_error", message="Неизвестный формат: " + ", ".join(unknown))
        return 2
    formats = requested
    exports = output_root / "exports"
    exports.mkdir(parents=True, exist_ok=True)
    markdown_path = find_file(result, "_enhanced.md") or find_file(result, ".md", "_enhanced")
    exact_json = result / f"{result.name}.json"
    json_candidates = [p for p in result.rglob("*.json") if "_model" not in p.name and "manifest" not in p.name]
    json_path = exact_json if exact_json.exists() else (sorted(json_candidates, key=lambda p: len(p.parts))[0] if json_candidates else None)
    if not markdown_path or not json_path:
        emit("export_error", message="Markdown или JSON результата не найден")
        return 2
    base = result.name
    markdown = markdown_path.read_text(encoding="utf-8", errors="replace")
    total = len(formats)
    done = 0
    emit("export_started", total=total, formats=formats)
    for format_name in formats:
        emit("export_format_started", format=format_name, done=done, total=total)
        path: Path | None = None
        if format_name == "md":
            path = exports / f"{base}.md"
            path.write_text(rewrite_image_links(markdown, result, path), encoding="utf-8")
        elif format_name == "split_md":
            path = exports / f"{base}_chapters"
            make_split_md(markdown, path)
            for chapter in path.glob("*.md"):
                chapter.write_text(rewrite_image_links(chapter.read_text(encoding="utf-8"), result, chapter), encoding="utf-8")
        elif format_name == "html":
            path = exports / f"{base}.html"
            path.write_text(markdown_to_html(markdown, base), encoding="utf-8")
        elif format_name == "docx":
            path = exports / f"{base}.docx"
            add_markdown_to_docx(markdown, path, base, markdown_path.parent)
        elif format_name == "searchable_pdf":
            path = exports / f"{base}_searchable.pdf"
            format_done = done
            format_total = total

            def on_page(pages_done: int, pages_total: int, _done: int = format_done, _total: int = format_total) -> None:
                emit(
                    "searchable_pdf_progress",
                    pages_done=pages_done,
                    pages_total=pages_total,
                    format_done=_done,
                    format_total=_total,
                )

            try:
                kwargs = {
                    "force_grayscale": args.force_grayscale,
                    "overlay_crops": args.overlay_crops,
                    "result_dir": result,
                    "progress": on_page,
                }
                if args.downsample_dpi is not None:
                    kwargs["downsample_dpi"] = args.downsample_dpi
                kwargs["jpeg_quality"] = args.jpeg_quality
                make_searchable_pdf(source, json_path, path, **kwargs)
            except Exception as exc:
                emit("export_error", format=format_name, message=f"Searchable PDF: {exc}")
                return 3
            if not path.exists():
                path = None
        elif format_name == "clean_pdf":
            path = exports / f"{base}_clean.pdf"
            process = subprocess.Popen([
                sys.executable, str(RECONSTRUCT_SCRIPT), "--source", str(source), "--json", str(json_path),
                "--result", str(result), "--output", str(path),
            ], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, bufsize=1)
            assert process.stdout is not None
            for line in process.stdout:
                match = re.match(r"CLEAN_PDF_PROGRESS\s+(\d+)\s+(\d+)", line.strip())
                if match:
                    emit("clean_pdf_progress", pages_done=int(match.group(1)), pages_total=int(match.group(2)),
                         format_done=done, format_total=total)
                elif line.strip():
                    emit("export_log", message=line.strip()[-1000:])
            if process.wait() != 0:
                emit("export_error", format=format_name, message="Не удалось собрать чистый PDF")
                return 3
        elif format_name == "zip":
            path = exports / f"{base}.zip"
            with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
                for item in output_root.rglob("*"):
                    if not item.is_file() or item == path or ".job" in item.parts:
                        continue
                    archive.write(item, item.relative_to(output_root))
        done += 1
        emit("export_progress", format=format_name, path=str(path) if path else "", done=done, total=total)
    emit("export_done", directory=str(exports))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
